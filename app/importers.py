"""Convert user order tables into an explicit, auditable scenario document."""

import csv
import json
from math import isfinite
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path

from .scenarios import ScenarioError, validate_document

ALIASES = {
    "id": ("order_id", "order", "po", "po_number"),
    "factory_id": ("factory", "factory_code"),
    "line_id": ("line", "sewing_line"),
    "style_id": ("style", "style_code"),
    "quantity": ("qty", "units", "order_quantity"),
    "due_date": ("deadline", "ship_date", "shipment_deadline"),
    "material_date": ("material_ready", "fabric_ready", "material_availability"),
    "smv": ("standard_minutes", "standard_minute_value"),
    "category": ("garment_category", "product_category"),
    "machine": ("required_machine",),
}


def _value(row, key, default=None):
    for name in (key, *ALIASES.get(key, ())):
        val = row.get(name)
        if val is not None and str(val).strip() != "":
            return val
    return default


def _day(value, label):
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    raw = str(value or "").strip()
    try:
        return date.fromisoformat(raw).isoformat()
    except ValueError:
        raise ScenarioError(f"{label} must be an ISO date (YYYY-MM-DD)") from None


def _number(value, label, default=None):
    if value is None or str(value).strip() == "":
        value = default
    try:
        number = float(value)
        if not isfinite(number):
            raise ValueError
        return number
    except (ValueError, TypeError, OverflowError):
        raise ScenarioError(f"{label} must be numeric") from None


def _integer(value, label, default=None):
    number = _number(value, label, default)
    if not number.is_integer():
        raise ScenarioError(f"{label} must be a whole number")
    return int(number)


def _boolean(value, default=True):
    if value is None or str(value).strip() == "":
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().lower()
    if text in ("true", "yes", "y", "1", "released", "pass"):
        return True
    if text in ("false", "no", "n", "0", "hold", "pending"):
        return False
    raise ScenarioError(f"Boolean value {value!r} must be yes/no or true/false")


def _records(rows, filename):
    if not isinstance(rows, list):
        raise ScenarioError("Order rows must be an array of objects")
    if not rows:
        raise ScenarioError("Order file has no data rows")
    required = ("id", "quantity", "due_date", "material_date")
    data = {
        key: []
        for key in (
            "factories",
            "lines",
            "styles",
            "orders",
            "materials",
            "production",
            "downtime",
            "quality_holds",
            "labor",
            "incidents",
            "actions",
        )
    }
    factories = {}
    lines = {}
    styles = {}
    assumed = set()
    material_dates = []
    declared_dates = set()
    source_labels = set()
    for index, raw in enumerate(rows, 2):
        if not isinstance(raw, dict):
            raise ScenarioError(f"Row {index} must be an object with order columns")
        row = {
            str(k).strip().lower().replace(" ", "_"): v
            for k, v in raw.items()
            if k is not None
        }
        if not any(v is not None and str(v).strip() for v in row.values()):
            continue
        for field in required:
            if _value(row, field) is None:
                raise ScenarioError(f"Row {index} needs {field}")
        order_id = str(_value(row, "id")).strip()
        factory_id = str(_value(row, "factory_id", "F1")).strip()
        line_id = str(_value(row, "line_id", "L1")).strip()
        style_id = str(_value(row, "style_id", "S1")).strip()
        if _value(row, "factory_id") is None:
            assumed.add("Missing factory assigned to F1")
        if _value(row, "line_id") is None:
            assumed.add("Missing line assigned to L1")
        if _value(row, "style_id") is None:
            assumed.add("Missing style assigned to S1")
        qty = _number(_value(row, "quantity"), f"Row {index} quantity")
        if qty < 1 or not qty.is_integer():
            raise ScenarioError(f"Row {index} quantity must be a positive whole number")
        due = _day(_value(row, "due_date"), f"Row {index} due_date")
        material = _day(_value(row, "material_date"), f"Row {index} material_date")
        material_dates.append(material)
        if _value(row, "as_of") is not None:
            declared_dates.add(_day(_value(row, "as_of"), f"Row {index} as_of"))
        if _value(row, "synthetic") is not None:
            source_labels.add(_boolean(_value(row, "synthetic")))
        category = str(_value(row, "category", "woven")).strip().lower()
        machine = str(_value(row, "machine", "lockstitch")).strip().lower()
        smv = _number(_value(row, "smv"), f"Row {index} smv", 20)
        if smv <= 0:
            raise ScenarioError(f"Row {index} smv must be positive")
        if _value(row, "category") is None:
            assumed.add("Missing garment category treated as woven")
        if _value(row, "machine") is None:
            assumed.add("Missing machine treated as lockstitch")
        if _value(row, "smv") is None:
            assumed.add("Missing SMV treated as 20 minutes")
        if factory_id not in factories:
            factory = dict(
                id=factory_id,
                name=str(_value(row, "factory_name", factory_id)),
                region=str(_value(row, "region", "Unspecified")),
                daily_cut=_integer(_value(row, "daily_cut"), "daily_cut", 2000),
                daily_finish=_integer(
                    _value(row, "daily_finish"), "daily_finish", 2000
                ),
                daily_qc=_integer(_value(row, "daily_qc"), "daily_qc", 2000),
                daily_pack=_integer(_value(row, "daily_pack"), "daily_pack", 2000),
                cost=_number(_value(row, "factory_cost"), "factory_cost", 5),
                quality=_number(
                    _value(row, "factory_quality"), "factory_quality", 0.95
                ),
                reliability=_number(
                    _value(row, "factory_reliability"), "factory_reliability", 0.85
                ),
                logistics_days=_integer(
                    _value(row, "logistics_days"), "logistics_days", 2
                ),
                compliance=_boolean(_value(row, "compliance")),
                capabilities=[],
            )
            factories[factory_id] = factory
            if _value(row, "daily_cut") is None:
                assumed.add(
                    "Unprovided factory stage capacities treated as 2,000 units/day"
                )
            if _value(row, "factory_cost") is None:
                assumed.add(
                    "Unprovided factory cost, quality, and reliability use illustrative defaults"
                )
        # Repeated entities must not silently depend on row order.
        for column, field, convert in [
            ("factory_cost", "cost", _number),
            ("factory_quality", "quality", _number),
            ("factory_reliability", "reliability", _number),
            ("logistics_days", "logistics_days", _integer),
            *[
                (key, key, _integer)
                for key in ("daily_cut", "daily_finish", "daily_qc", "daily_pack")
            ],
        ]:
            supplied = _value(row, column)
            if (
                supplied is not None
                and convert(supplied, column) != factories[factory_id][field]
            ):
                raise ScenarioError(
                    f"Factory {factory_id} has conflicting {column} at row {index}"
                )
        if (
            _value(row, "compliance") is not None
            and _boolean(_value(row, "compliance"))
            != factories[factory_id]["compliance"]
        ):
            raise ScenarioError(
                f"Factory {factory_id} has conflicting compliance at row {index}"
            )
        if category not in factories[factory_id]["capabilities"]:
            factories[factory_id]["capabilities"].append(category)
        if line_id not in lines:
            lines[line_id] = dict(
                id=line_id,
                factory_id=factory_id,
                name=str(_value(row, "line_name", line_id)),
                operators=_integer(_value(row, "operators"), "operators", 30),
                efficiency=_number(_value(row, "efficiency"), "efficiency", 0.75),
                skills=[],
                machines=[],
                maintenance=[],
            )
            if _value(row, "operators") is None:
                assumed.add(
                    "Unprovided line staffing and efficiency use 30 operators at 75%"
                )
        elif lines[line_id]["factory_id"] != factory_id:
            raise ScenarioError(f"Line {line_id} is assigned to more than one factory")
        for column, convert in [("operators", _integer), ("efficiency", _number)]:
            supplied = _value(row, column)
            if (
                supplied is not None
                and convert(supplied, column) != lines[line_id][column]
            ):
                raise ScenarioError(
                    f"Line {line_id} has conflicting {column} at row {index}"
                )
        if category not in lines[line_id]["skills"]:
            lines[line_id]["skills"].append(category)
        if machine not in lines[line_id]["machines"]:
            lines[line_id]["machines"].append(machine)
        if style_id not in styles:
            styles[style_id] = dict(
                id=style_id,
                name=str(_value(row, "style_name", style_id)),
                category=category,
                smv=smv,
                bottleneck=str(_value(row, "bottleneck", "Sewing")),
                machine=machine,
                operations=[["Sewing", smv]],
            )
        elif (
            styles[style_id]["category"],
            styles[style_id]["machine"],
            styles[style_id]["smv"],
        ) != (category, machine, smv):
            raise ScenarioError(
                f"Style {style_id} has conflicting category, machine, or SMV"
            )
        target = round(
            _number(
                _value(row, "target_cost"),
                "target_cost",
                qty * factories[factory_id]["cost"],
            ),
            2,
        )
        data["orders"].append(
            dict(
                id=order_id,
                customer=str(_value(row, "customer", "Unspecified")),
                style_id=style_id,
                factory_id=factory_id,
                line_id=line_id,
                quantity=int(qty),
                due_date=due,
                material_date=material,
                target_cost=target,
                priority=_integer(_value(row, "priority"), "priority", 1),
                margin_per_unit=_number(
                    _value(row, "margin_per_unit"), "margin_per_unit", 3
                ),
                min_lot=_integer(_value(row, "min_lot"), "min_lot", 1),
                quality_release=_boolean(_value(row, "quality_release")),
                status="active",
            )
        )
        data["materials"].append(
            dict(
                order_id=order_id,
                fabric_ready=_day(
                    _value(row, "fabric_ready", material), "fabric_ready"
                ),
                trims_ready=_day(_value(row, "trims_ready", material), "trims_ready"),
                shortage_units=_integer(
                    _value(row, "shortage_units"), "shortage_units", 0
                ),
            )
        )
    data["factories"] = list(factories.values())
    data["lines"] = list(lines.values())
    data["styles"] = list(styles.values())
    if not data["orders"]:
        raise ScenarioError("Order file has no data rows")
    if len(declared_dates) > 1:
        raise ScenarioError("Order rows have conflicting as_of dates")
    if len(source_labels) > 1:
        raise ScenarioError("Order rows have conflicting synthetic labels")
    if not declared_dates:
        assumed.add("Scenario date inferred from the earliest material date")
    assumed.add(
        "Historical production, downtime, quality and incident metrics remain unavailable unless supplied in a full scenario JSON"
    )
    assumed.add(
        "Factory capabilities and line skills/machines are inferred from assigned orders"
    )
    assumed.add(
        "Each style uses one sewing operation at its supplied or default SMV; detailed operation routes need a full scenario JSON"
    )
    assumed.add(
        "No holidays are assumed for an order table; supply meta.holidays in a full scenario JSON"
    )
    meta = dict(
        id="imported",
        name=Path(filename).stem.replace("_", " ").replace("-", " ").title(),
        description=f"{len(data['orders'])} orders imported from {Path(filename).name}",
        synthetic=next(iter(source_labels)) if source_labels else False,
        as_of=next(iter(declared_dates)) if declared_dates else min(material_dates),
        source_format=Path(filename).suffix.lower().lstrip("."),
        assumptions=sorted(assumed),
    )
    return validate_document(dict(schema_version=1, meta=meta, data=data))


def import_file(payload, filename):
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        try:
            rows = list(csv.DictReader(StringIO(payload.decode("utf-8-sig"))))
        except UnicodeDecodeError:
            raise ScenarioError("CSV must be UTF-8 encoded") from None
        return _records(rows, filename)
    if extension == ".xlsx":
        from openpyxl import load_workbook

        try:
            book = load_workbook(BytesIO(payload), read_only=True, data_only=True)
            sheet = book["Orders"] if "Orders" in book.sheetnames else book.active
            values = sheet.iter_rows(values_only=True)
            headings = next(values, None)
            rows = (
                [dict(zip(headings, values_row)) for values_row in values]
                if headings
                else []
            )
            book.close()
        except Exception as exc:
            raise ScenarioError(f"Could not read Excel workbook: {exc}") from None
        return _records(rows, filename)
    if extension == ".json":
        try:
            source = json.loads(payload.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            raise ScenarioError("Upload a valid UTF-8 JSON file") from None
        if isinstance(source, list):
            return _records(source, filename)
        if isinstance(source, dict) and "orders" in source and "data" not in source:
            return _records(source["orders"], filename)
        return validate_document(source)
    raise ScenarioError("Supported files are .json, .csv, and .xlsx")
